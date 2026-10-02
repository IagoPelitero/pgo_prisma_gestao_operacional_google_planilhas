#!/usr/bin/env python3
"""
============================================================================
PGO — conferir-excel.py · o Excel volta a ser código que CARREGA?
============================================================================
POR QUE ISTO EXISTE. O PO instala o PGO copiando as abas desta planilha para
o Apps Script — é o transporte dele, não só material de leitura. E uma
planilha chegou até ele com `Indicadores.gs` contendo um bloco de comentário
sem a abertura `/**`. No Apps Script os arquivos .gs vivem num escopo global
só: UM erro de sintaxe em um arquivo impede o projeto INTEIRO de carregar.
Nem `instalarRECC()` roda. A tela não abre.

Ninguém percebeu porque ninguém tinha lido a planilha DE VOLTA. O gerador
escrevia, e a conferência olhava o repositório — nunca o que foi gerado.

O que este conferidor faz, aba por aba:

  1. extrai o código da aba, tirando a marca `* ` do número da linha;
  2. confere que ele é IGUAL ao arquivo do repositório, caractere a caractere;
  3. confere que a numeração é contígua, sem linha perdida no caminho.

E no fim roda os .gs juntos, no Node, como o Apps Script faria.

Sai com código 1 se qualquer aba falhar — é portão, não relatório.
============================================================================
"""

import io
import os
import re
import subprocess
import sys
import tempfile

import openpyxl

RAIZ = os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
PLANILHA = os.path.join(RAIZ, 'Evolucao', 'pacote', 'PGO-codigo-completo.xlsx')

# De qual pasta vem cada arquivo. É a mesma lista que o gerador usa.
ONDE_MORA = {'.gs': 'Back-End', '.html': 'Front-End'}
FORA_DO_PADRAO = {'conferir-projeto.gs': 'Evolucao'}


def codigo_da_aba(aba):
    """O código de uma aba, e os números de linha que ela declara."""
    linhas = []
    numeros = []
    for r in range(2, aba.max_row + 1):
        marca = aba.cell(row=r, column=1).value
        codigo = aba.cell(row=r, column=2).value
        if marca is None and codigo is None:
            continue
        texto = str(marca).strip() if marca is not None else ''
        # A linha editada vem marcada: "* 169" em vez de 169.
        achado = re.match(r'^\*?\s*(\d+)$', texto)
        if achado:
            numeros.append(int(achado.group(1)))
        linhas.append('' if codigo is None else str(codigo))
    return linhas, numeros


def caminho_do_arquivo(nome):
    if nome in FORA_DO_PADRAO:
        return os.path.join(RAIZ, FORA_DO_PADRAO[nome], nome)
    extensao = os.path.splitext(nome)[1]
    return os.path.join(RAIZ, ONDE_MORA[extensao], nome)


def main():
    if not os.path.exists(PLANILHA):
        print('A planilha não existe. Rode gerar-excel.py antes.')
        return 1

    livro = openpyxl.load_workbook(PLANILHA)
    problemas = []
    conferidas = 0
    pasta = tempfile.mkdtemp(prefix='pgo-excel-')

    print('\nO EXCEL VOLTA A SER CÓDIGO?')
    print('-' * 72)

    for titulo in livro.sheetnames:
        if titulo == 'LEIA-ME':
            continue

        # Toda aba é um arquivo-fonte. O losango só marca o que mudou nesta
        # rodada — ele não faz parte do nome do arquivo.
        nome = titulo.replace('◆ ', '')
        aba = livro[titulo]
        linhas, numeros = codigo_da_aba(aba)
        # SEM um '\n' a mais no fim: o gerador já grava a linha vazia final do
        # arquivo como uma linha da aba, então juntar com '\n' reconstrói o
        # arquivo exatamente. Somar outro acusaria as 21 abas de diferentes.
        doExcel = '\n'.join(linhas)

        # 1. a numeração não pode pular: linha perdida é dado perdido.
        esperado = list(range(1, len(numeros) + 1))
        if numeros != esperado:
            faltando = sorted(set(esperado) - set(numeros))
            problemas.append(nome + ': a numeração pula — faltam as linhas '
                             + str(faltando[:10]))

        # 2. tem de bater com o arquivo do repositório, caractere a caractere.
        caminho = caminho_do_arquivo(nome)
        doRepositorio = io.open(caminho, encoding='utf-8').read()
        if doExcel != doRepositorio:
            daPlanilha = doExcel.split('\n')
            doDisco = doRepositorio.split('\n')
            onde = 'tamanhos diferentes (%d x %d linhas)' % (len(daPlanilha), len(doDisco))
            for i in range(min(len(daPlanilha), len(doDisco))):
                if daPlanilha[i] != doDisco[i]:
                    onde = 'primeira diferença na linha %d' % (i + 1)
                    break
            problemas.append(nome + ': o que está na aba não é o arquivo — ' + onde)

        if nome.endswith('.gs'):
            with io.open(os.path.join(pasta, nome), 'w', encoding='utf-8') as f:
                f.write(doExcel)

        conferidas += 1
        print('  %-24s %5d linhas' % (nome, len(linhas)))

    # 3. e os .gs juntos têm de CARREGAR, como no Apps Script.
    print('-' * 72)
    juntos = os.path.join(pasta, '__todos.js')
    with io.open(juntos, 'w', encoding='utf-8') as f:
        for nome in sorted(os.listdir(pasta)):
            if nome.endswith('.gs') and nome != 'conferir-projeto.gs':
                f.write(io.open(os.path.join(pasta, nome), encoding='utf-8').read())
                f.write('\n')
    conferencia = subprocess.run(
        ['node', '--check', juntos], capture_output=True, text=True)
    if conferencia.returncode != 0:
        primeira = (conferencia.stderr or '').strip().split('\n')
        problemas.append('os .gs do Excel NÃO carregam juntos: '
                         + ' / '.join(primeira[:4]))
    else:
        print('  os .gs do Excel carregam juntos, como no Apps Script.')

    print('-' * 72)
    if problemas:
        print('%d aba(s) conferida(s) — %d PROBLEMA(S):\n' % (conferidas, len(problemas)))
        for problema in problemas:
            print('  ' + problema)
        return 1

    print('%d abas conferidas, nenhuma diferença. A planilha serve para colar.'
          % conferidas)
    return 0


if __name__ == '__main__':
    sys.exit(main())
