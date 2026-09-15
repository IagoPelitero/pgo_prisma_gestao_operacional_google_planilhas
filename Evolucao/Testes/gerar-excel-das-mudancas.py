# -*- coding: utf-8 -*-
"""
============================================================================
PGO — gerar-excel-das-mudancas.py · só o que mudou, para colar no Apps Script
============================================================================
    python3 Evolucao/Testes/gerar-excel-das-mudancas.py [referencia]

O `gerar-excel.py` escreve o sistema INTEIRO — 21 abas, 18 mil linhas. Serve
para quem instala do zero. Não serve para quem já tem o PGO rodando e só
precisa atualizar: ali, a pergunta é "quais arquivos eu recolo?", e uma
planilha com tudo obriga a recolar tudo.

Esta aqui responde essa pergunta. Compara o estado atual com uma referência
do git (por padrão, o último commit) e escreve uma aba por arquivo ALTERADO,
mais uma capa dizendo o que mudou em cada um.

Só entra o que vai para o Apps Script: Back-End e Front-End. A pasta Evolucao
— documentação, testes, ferramentas — não é copiada para lá, e listá-la aqui
mandaria alguém colar o que não tem onde morar.
============================================================================
"""
import io
import os
import subprocess
import sys

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

RAIZ = os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
PASTAS_DO_APPS_SCRIPT = ('Back-End', 'Front-End')
DESTINO = os.path.join(RAIZ, 'Evolucao', 'pacote', 'PGO-mudancas.xlsx')

# O Excel corta nome de aba em 31 caracteres e não aceita alguns sinais.
LIMITE_DO_NOME_DA_ABA = 31


def rodar(*argumentos):
    return subprocess.run(argumentos, cwd=RAIZ, capture_output=True,
                          text=True, check=True).stdout


def arquivos_que_mudaram(referencia):
    """Os arquivos do Apps Script alterados em relação à referência."""
    saida = rodar('git', 'status', '--porcelain')
    mudados = set()
    for linha in saida.split('\n'):
        if not linha.strip():
            continue
        caminho = linha[3:].strip().strip('"')
        # renomeado vem como "antigo -> novo"
        if ' -> ' in caminho:
            caminho = caminho.split(' -> ')[1]
        if caminho.startswith(PASTAS_DO_APPS_SCRIPT):
            mudados.add(caminho)

    # E o que já foi commitado depois da referência, quando ela não é HEAD.
    if referencia != 'HEAD':
        for caminho in rodar('git', 'diff', '--name-only', referencia).split('\n'):
            caminho = caminho.strip()
            if caminho.startswith(PASTAS_DO_APPS_SCRIPT):
                mudados.add(caminho)

    return sorted(mudados)


def resumo_da_mudanca(caminho, referencia):
    """Quantas linhas entraram e saíram, para a capa."""
    try:
        saida = rodar('git', 'diff', '--numstat', referencia, '--', caminho)
    except subprocess.CalledProcessError:
        return ('novo', 0, 0)
    linha = saida.strip().split('\n')[0] if saida.strip() else ''
    if not linha:
        return ('novo', 0, 0)
    partes = linha.split('\t')
    if len(partes) < 2:
        return ('alterado', 0, 0)
    entraram = 0 if partes[0] == '-' else int(partes[0])
    sairam = 0 if partes[1] == '-' else int(partes[1])
    return ('alterado', entraram, sairam)


def nome_da_aba(caminho, jaUsados):
    base = os.path.basename(caminho)
    nome = base.replace('.', '_')[:LIMITE_DO_NOME_DA_ABA]
    contador = 2
    while nome in jaUsados:
        sufixo = '_' + str(contador)
        nome = (base.replace('.', '_')[:LIMITE_DO_NOME_DA_ABA - len(sufixo)]) + sufixo
        contador += 1
    jaUsados.add(nome)
    return nome


def escrever_capa(planilha, arquivos, referencia):
    aba = planilha.active
    aba.title = 'LEIA-ME'

    titulo = Font(name='Arial', size=14, bold=True)
    cabecalho = Font(name='Arial', size=10, bold=True, color='FFFFFF')
    corpo = Font(name='Arial', size=10)
    fundo = PatternFill('solid', fgColor='1F3864')

    aba['A1'] = 'PGO — só os arquivos que mudaram'
    aba['A1'].font = titulo
    aba['A3'] = ('Compara o estado atual com: ' + referencia)
    aba['A3'].font = corpo
    aba['A4'] = ('Recole no Apps Script apenas os arquivos listados abaixo. '
                 'Cada um tem a sua aba, com o conteúdo inteiro, uma linha de '
                 'código por linha da planilha.')
    aba['A4'].font = corpo
    aba['A4'].alignment = Alignment(wrap_text=True, vertical='top')
    aba.merge_cells('A4:D4')
    aba.row_dimensions[4].height = 42

    aba['A6'] = 'ATENÇÃO AO NOME: no Apps Script o arquivo se chama Base (sem ".gs") '
    aba['A6'].font = Font(name='Arial', size=10, bold=True)

    linha = 8
    for coluna, texto in enumerate(['Arquivo', 'No Apps Script', 'Linhas',
                                    'Mudança'], start=1):
        celula = aba.cell(row=linha, column=coluna, value=texto)
        celula.font = cabecalho
        celula.fill = fundo

    for caminho in arquivos:
        linha += 1
        base = os.path.basename(caminho)
        com_texto = io.open(os.path.join(RAIZ, caminho), encoding='utf-8').read()
        estado, entraram, sairam = resumo_da_mudanca(caminho, referencia)
        mudanca = ('arquivo novo' if estado == 'novo'
                   else '+' + str(entraram) + ' / -' + str(sairam) + ' linhas')
        valores = [caminho, base.rsplit('.', 1)[0],
                   len(com_texto.split('\n')), mudanca]
        for coluna, valor in enumerate(valores, start=1):
            celula = aba.cell(row=linha, column=coluna, value=valor)
            celula.font = corpo

    for coluna, largura in enumerate([34, 22, 10, 20], start=1):
        aba.column_dimensions[get_column_letter(coluna)].width = largura


def escrever_arquivo(planilha, caminho, jaUsados):
    texto = io.open(os.path.join(RAIZ, caminho), encoding='utf-8').read()
    linhas = texto.split('\n')

    aba = planilha.create_sheet(nome_da_aba(caminho, jaUsados))
    aba['A1'] = 'nº'
    aba['B1'] = caminho
    for celula in (aba['A1'], aba['B1']):
        celula.font = Font(name='Arial', size=10, bold=True, color='FFFFFF')
        celula.fill = PatternFill('solid', fgColor='1F3864')

    monoespacada = Font(name='Consolas', size=9)
    for i, conteudo in enumerate(linhas):
        aba.cell(row=i + 2, column=1, value=i + 1).font = Font(name='Arial', size=8)
        # Uma linha que comece com "=" seria lida como fórmula pelo Excel. O
        # tipo da célula fica em texto para o código chegar inteiro do outro lado.
        celula = aba.cell(row=i + 2, column=2)
        celula.value = conteudo
        celula.data_type = 's'
        celula.font = monoespacada

    aba.column_dimensions['A'].width = 6
    aba.column_dimensions['B'].width = 120
    aba.freeze_panes = 'B2'
    return len(linhas)


def main():
    referencia = sys.argv[1] if len(sys.argv) > 1 else 'HEAD'
    arquivos = arquivos_que_mudaram(referencia)

    if not arquivos:
        print('Nenhum arquivo do Apps Script mudou em relação a ' + referencia + '.')
        return

    planilha = Workbook()
    escrever_capa(planilha, arquivos, referencia)

    jaUsados = set()
    total = 0
    for caminho in arquivos:
        total += escrever_arquivo(planilha, caminho, jaUsados)

    planilha.save(DESTINO)
    print('Planilha das mudanças em ' + DESTINO)
    print('  ' + str(len(arquivos)) + ' arquivos alterados, '
          + format(total, ',d').replace(',', '.') + ' linhas')
    for caminho in arquivos:
        print('    ' + caminho)


main()
