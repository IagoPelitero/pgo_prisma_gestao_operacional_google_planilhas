# -*- coding: utf-8 -*-
"""
============================================================================
PGO — gerar-excel.py · o código inteiro numa planilha, um arquivo por aba
============================================================================
    python3 Evolucao/Testes/gerar-excel.py

Gera `Evolucao/pacote/PGO-codigo-completo.xlsx`: uma aba por arquivo, com o
número da linha ao lado do código, para conferir e copiar de onde não dá para
clonar o repositório.

É a única ferramenta do projeto em Python — openpyxl é o que sabe escrever
.xlsx, e vale mais usar a biblioteca certa do que reimplementar o formato.

----------------------------------------------------------------------------
UM AVISO QUE PRECISA ESTAR AQUI E NA PRIMEIRA ABA
----------------------------------------------------------------------------
Excel NÃO é bom transporte de código. Copiar uma coluna de volta para o Apps
Script costuma trazer o que o Excel achou que era melhor: aspas retas viram
curvas, uma linha em branco some, o recuo se perde. Para INSTALAR, o caminho
é `Evolucao/pacote/` — três arquivos de texto, prontos para colar.

Esta planilha serve para LER, revisar e comparar. É para isso que ela é boa.
============================================================================
"""

import os
import subprocess
import sys

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

RAIZ = os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
DESTINO = os.path.join(RAIZ, 'Evolucao', 'pacote', 'PGO-codigo-completo.xlsx')

# As cores do tema padrão do próprio sistema, para a planilha parecer dele.
AZUL = '0B77CE'
AZUL_CLARO = 'E8F1FB'
CINZA = '59687F'
LINHA = 'D8E1ED'

FONTE_CODIGO = Font(name='Consolas', size=10)
FONTE_NUMERO = Font(name='Consolas', size=9, color=CINZA)
FONTE_TITULO = Font(name='Arial', size=11, bold=True, color='FFFFFF')
BORDA_BAIXO = Border(bottom=Side(style='thin', color=LINHA))


def arquivos_do_codigo():
    """Os arquivos que a pessoa precisa ter no Apps Script, em ordem de leitura."""
    grupos = [
        ('Back-End', '.gs', 'servidor'),
        ('Front-End', '.html', 'tela'),
    ]
    achados = []
    for pasta, extensao, papel in grupos:
        caminho = os.path.join(RAIZ, pasta)
        for nome in sorted(os.listdir(caminho)):
            if nome.endswith(extensao):
                achados.append((pasta, nome, papel))
    # A conferência avulsa vai junto: é o que a pessoa roda quando algo falta.
    achados.append(('Evolucao', 'conferir-projeto.gs', 'conferência'))
    return achados


def escrever_aba_do_arquivo(planilha, pasta, nome, papel):
    aba = planilha.create_sheet(title=nome[:31])

    aba['A1'] = 'Linha'
    aba['B1'] = pasta + '/' + nome
    for celula in (aba['A1'], aba['B1']):
        celula.font = FONTE_TITULO
        celula.fill = PatternFill('solid', fgColor=AZUL)
        celula.alignment = Alignment(vertical='center')
    aba.row_dimensions[1].height = 22

    caminho = os.path.join(RAIZ, pasta, nome)
    with open(caminho, encoding='utf-8') as arquivo:
        linhas = arquivo.read().split('\n')

    for numero, conteudo in enumerate(linhas, start=1):
        alvo = aba.cell(row=numero + 1, column=1, value=numero)
        alvo.font = FONTE_NUMERO
        alvo.alignment = Alignment(horizontal='right')

        codigo = aba.cell(row=numero + 1, column=2)
        codigo.value = conteudo
        # Uma linha que comece com "=" viraria fórmula. Não há nenhuma hoje,
        # e a trava fica para o dia em que houver.
        codigo.data_type = 's'
        codigo.font = FONTE_CODIGO
        codigo.alignment = Alignment(vertical='top')

    aba.column_dimensions['A'].width = 7
    aba.column_dimensions['B'].width = 120
    aba.freeze_panes = 'A2'
    return len(linhas)


def escrever_leia_me(planilha, inventario):
    aba = planilha.create_sheet(title='LEIA-ME', index=0)
    aba.column_dimensions['A'].width = 30
    aba.column_dimensions['B'].width = 14
    aba.column_dimensions['C'].width = 78

    linhas = [
        ('PGO — Prisma Gestão Operacional', '', ''),
        ('Operação RECC · Porto Seguro', '', ''),
        ('', '', ''),
        ('PARA QUE SERVE ESTA PLANILHA', '', ''),
        ('', '', 'Ler, revisar e comparar o código. Uma aba por arquivo, com o '
                 'número da linha ao lado.'),
        ('', '', ''),
        ('PARA INSTALAR, NÃO USE ESTA PLANILHA', '', ''),
        ('', '', 'Excel não é bom transporte de código: copiar uma coluna de volta '
                 'costuma trazer o que o Excel'),
        ('', '', 'achou melhor — aspas retas viram curvas, linha em branco some, o '
                 'recuo se perde.'),
        ('', '', ''),
        ('', '', 'Use a pasta Evolucao/pacote/ do repositório: são TRÊS arquivos de '
                 'texto, prontos para colar.'),
        ('', '', '   Codigo.gs       ->  arquivo .gs chamado "Codigo"'),
        ('', '', '   Index.html      ->  arquivo HTML chamado "Index"'),
        ('', '', '   SemAcesso.html  ->  arquivo HTML chamado "SemAcesso"'),
        ('', '', ''),
        ('', '', 'No Apps Script o arquivo se chama "Index", e não "Index.html": sem '
                 'extensão, sem acento,'),
        ('', '', 'com as maiúsculas iguais.'),
        ('', '', ''),
        ('DEPOIS DE INSTALAR', '', ''),
        ('', '', 'instalarRECC()        cria as 13 abas e cadastra você como o '
                 'primeiro administrador'),
        ('', '', 'diagnosticoRECC()     confere o sistema inteiro e diz o que falta'),
        ('', '', 'oQueFaltaNoProjeto()  de conferir-projeto.gs: diz qual arquivo '
                 'ficou para trás na cópia'),
        ('', '', ''),
        ('OS ARQUIVOS', '', ''),
    ]

    for texto_a, texto_b, texto_c in linhas:
        aba.append([texto_a, texto_b, texto_c])

    cabecalho = aba.max_row + 1
    aba.cell(row=cabecalho, column=1, value='Arquivo')
    aba.cell(row=cabecalho, column=2, value='Linhas')
    aba.cell(row=cabecalho, column=3, value='Onde fica / o que é')
    for coluna in range(1, 4):
        celula = aba.cell(row=cabecalho, column=coluna)
        celula.font = FONTE_TITULO
        celula.fill = PatternFill('solid', fgColor=AZUL)

    for pasta, nome, papel, quantas in inventario:
        aba.append([nome, quantas, pasta + '/  ·  ' + papel])

    for linha in aba.iter_rows(min_row=1, max_row=aba.max_row):
        for celula in linha:
            if celula.font.color is None or celula.font.color.rgb != 'FFFFFFFF':
                if celula.row < cabecalho:
                    ehTitulo = celula.column == 1 and celula.value
                    celula.font = Font(name='Arial', size=11,
                                       bold=bool(ehTitulo), color=CINZA if not ehTitulo else '141F33')
                else:
                    celula.font = Font(name='Arial', size=10)
            celula.alignment = Alignment(vertical='center')

    aba.cell(row=1, column=1).font = Font(name='Arial', size=16, bold=True, color=AZUL)
    aba.cell(row=2, column=1).font = Font(name='Arial', size=11, color=CINZA)
    aba.freeze_panes = 'A' + str(cabecalho + 1)
    return aba


def gerar():
    planilha = Workbook()
    planilha.remove(planilha.active)

    inventario = []
    for pasta, nome, papel in arquivos_do_codigo():
        quantas = escrever_aba_do_arquivo(planilha, pasta, nome, papel)
        inventario.append((pasta, nome, papel, quantas))

    escrever_leia_me(planilha, inventario)

    os.makedirs(os.path.dirname(DESTINO), exist_ok=True)
    planilha.save(DESTINO)
    return inventario


if __name__ == '__main__':
    inventario = gerar()
    total = sum(item[3] for item in inventario)
    print('')
    print('Planilha gerada em ' + DESTINO)
    print('  ' + str(len(inventario) + 1) + ' abas  (LEIA-ME + '
          + str(len(inventario)) + ' arquivos)')
    print('  ' + '{:,}'.format(total).replace(',', '.') + ' linhas de código')
    print('  ' + str(round(os.path.getsize(DESTINO) / 1024)) + ' KB')
    print('')
